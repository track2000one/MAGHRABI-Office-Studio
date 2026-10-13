"""Premium, editable A4 Word output, with native heading and RTL support.

All visual elements use Word paragraphs and tables, not flattened screenshots.
Do not insert invented organizations, approval statements, dates, references, or statistics.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .content import Content

STYLE_NAMES = ("luxury", "corporate", "minimal")
DOCUMENT_TYPES = ("general", "report", "study", "strategy", "letter")
DOCUMENT_LABELS = {
    "general": ("مستند احترافي", "PROFESSIONAL DOCUMENT"),
    "report": ("تقرير رسمي", "OFFICIAL REPORT"),
    "study": ("دراسة تفصيلية", "DETAILED STUDY"),
    "strategy": ("خطة استراتيجية", "STRATEGIC PLAN"),
    "letter": ("خطاب إداري", "OFFICIAL LETTER"),
}


@dataclass(frozen=True)
class Theme:
    dark: str
    accent: str
    secondary: str
    muted: str
    pale: str
    rule: str
    face: str


THEMES = {
    "luxury": Theme("192B3E", "B28B50", "30465C", "607083", "F5F1E9", "DDD3C3", "Arial"),
    "corporate": Theme("163A52", "118D91", "345668", "607987", "EAF4F5", "D3E4E9", "Arial"),
    "minimal": Theme("28283C", "7665BA", "4B4B68", "74748A", "F2EFFA", "E2DDF0", "Arial"),
}


def _color(hex_value: str):
    return RGBColor.from_string(hex_value)


def _set_font(style, theme: Theme, size: int, color: str, *, bold=False):
    style.font.name = theme.face
    style.font.size = Pt(size)
    style.font.color.rgb = _color(color)
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn("w:" + key), theme.face)


def _paragraph_direction(paragraph, rtl: bool):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    properties = paragraph._p.get_or_add_pPr()
    direction = properties.find(qn("w:bidi"))
    if direction is None:
        direction = OxmlElement("w:bidi")
        properties.append(direction)
    direction.set(qn("w:val"), "1" if rtl else "0")
    if rtl:
        # Logical start is supported by Word and LibreOffice for RTL text.
        justification = properties.find(qn("w:jc"))
        if justification is not None:
            justification.set(qn("w:val"), "start")
    for run in paragraph.runs:
        run.font.complex_script = rtl
        fonts = run._r.get_or_add_rPr().rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run._r.get_or_add_rPr().insert(0, fonts)
        fonts.set(qn("w:cs"), "Arial")
    return paragraph


def _rule(paragraph, color: str, *, edge="bottom", size=13):
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    item = borders.find(qn("w:" + edge))
    if item is None:
        item = OxmlElement("w:" + edge)
        borders.append(item)
    item.set(qn("w:val"), "single")
    item.set(qn("w:sz"), str(size))
    item.set(qn("w:space"), "8")
    item.set(qn("w:color"), color)


def _shade_cell(cell, color: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, *, top=190, start=260, bottom=190, end=260):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _field(paragraph, instruction: str):
    run = paragraph.add_run()
    opening = OxmlElement("w:fldChar")
    opening.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    closing = OxmlElement("w:fldChar")
    closing.set(qn("w:fldCharType"), "end")
    for node in (opening, text, separator, display, closing):
        run._r.append(node)


def _bookmark(paragraph, number: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), f"maghrabi_section_{number}")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _contents_link(paragraph, text: str, number: int, theme: Theme):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), f"maghrabi_section_{number}")
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), theme.secondary)
    props.append(color)
    link_font = OxmlElement("w:rFonts")
    link_font.set(qn("w:cs"), theme.face)
    link_font.set(qn("w:ascii"), theme.face)
    link_font.set(qn("w:hAnsi"), theme.face)
    props.append(link_font)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def _add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _setup_styles(doc, theme: Theme, rtl: bool):
    normal = doc.styles["Normal"]
    _set_font(normal, theme, 11, theme.dark)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.48
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    _set_font(title, theme, 29, theme.dark, bold=True)
    title.paragraph_format.space_after = Pt(20)
    title.paragraph_format.line_spacing = 1.14
    title.paragraph_format.keep_with_next = True

    for key, size in (("Heading 1", 17), ("Heading 2", 13)):
        style = doc.styles[key]
        _set_font(style, theme, size, theme.dark, bold=True)
        style.paragraph_format.space_before = Pt(17)
        style.paragraph_format.space_after = Pt(11)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for label in ("Subtitle",):
        _set_font(doc.styles[label], theme, 13, theme.muted)


def _cover(doc, content: Content, theme: Theme, rtl: bool, document_type: str = "general"):
    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(18)
    lead.paragraph_format.space_after = Pt(14)
    run = lead.add_run("MAGHRABI  /  OFFICE STUDIO")
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = _color(theme.accent)
    _paragraph_direction(lead, rtl)
    _rule(lead, theme.accent, size=21)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(38)
    spacer.paragraph_format.space_after = Pt(0)

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(16)
    small = label.add_run(DOCUMENT_LABELS[document_type][0 if rtl else 1])
    small.bold = True
    small.font.size = Pt(11)
    small.font.color.rgb = _color(theme.accent)
    _paragraph_direction(label, rtl)

    main_title = doc.add_heading(content.title, 0)
    main_title.style = "Title"
    if len(content.title) > 90:
        for run in main_title.runs:
            run.font.size = Pt(23 if len(content.title) <= 140 else 20)
    _paragraph_direction(main_title, rtl)

    if content.summary:
        teaser = content.summary if len(content.summary) <= 175 else (
            content.summary[:172].rsplit(" ", 1)[0] + "…"
        )
        description = doc.add_paragraph(teaser)
        description.style = "Subtitle"
        description.paragraph_format.space_before = Pt(16)
        description.paragraph_format.space_after = Pt(18)
        description.paragraph_format.line_spacing = 1.55
        _paragraph_direction(description, rtl)

    panel = doc.add_table(rows=1, cols=1)
    panel.autofit = False
    panel.columns[0].width = Inches(6.5)
    cell = panel.cell(0, 0)
    _shade_cell(cell, theme.pale)
    _set_cell_margins(cell, top=245, bottom=245)
    message = cell.paragraphs[0]
    run = message.add_run(
        (f"المحاور الرئيسية  {len(content.sections):02d}" if rtl else
         f"PRIMARY SECTIONS  {len(content.sections):02d}")
    )
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = _color(theme.secondary)
    _paragraph_direction(message, rtl)

    _add_page_break(doc)


def _contents(doc, content: Content, theme: Theme, rtl: bool):
    heading = doc.add_heading("المحتويات" if rtl else "CONTENTS", 1)
    _paragraph_direction(heading, rtl)
    _rule(heading, theme.accent, size=17)

    instruction = doc.add_paragraph(
        "اضغط على عنوان المحور للانتقال إليه." if rtl else
        "Select a section title to navigate directly to it."
    )
    instruction.paragraph_format.space_after = Pt(16)
    for run in instruction.runs:
        run.font.color.rgb = _color(theme.muted)
        run.font.size = Pt(9)
    _paragraph_direction(instruction, rtl)

    for index, section in enumerate(content.sections, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(11)
        paragraph.paragraph_format.keep_together = True
        prefix = paragraph.add_run(f"{index:02d}     ")
        prefix.bold = True
        prefix.font.color.rgb = _color(theme.accent)
        _contents_link(paragraph, section.title, index, theme)
        _paragraph_direction(paragraph, rtl)
        _rule(paragraph, theme.rule, size=4)

    _add_page_break(doc)


def _summary(doc, content: Content, theme: Theme, rtl: bool, document_type: str = "general"):
    if not content.summary:
        return
    summary_heading = ("ملخص الدراسة" if rtl else "STUDY SUMMARY") if document_type == "study" else (
        ("ملخص الخطة" if rtl else "PLAN SUMMARY") if document_type == "strategy" else
        ("الملخص التنفيذي" if rtl else "EXECUTIVE SUMMARY")
    )
    p = doc.add_heading(summary_heading, 1)
    _paragraph_direction(p, rtl)
    _rule(p, theme.accent)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    _shade_cell(cell, theme.pale)
    _set_cell_margins(cell, top=210, bottom=210)
    body = cell.paragraphs[0]
    body.text = content.summary
    body.paragraph_format.line_spacing = 1.48
    body.paragraph_format.space_after = Pt(0)
    for run in body.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = _color(theme.dark)
    _paragraph_direction(body, rtl)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)


def _configure_header_footer(doc, theme: Theme, rtl: bool):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    mark = header.add_run("MAGHRABI   OFFICE STUDIO")
    mark.bold = True
    mark.font.size = Pt(8)
    mark.font.color.rgb = _color(theme.muted)
    _paragraph_direction(header, rtl)
    _rule(header, theme.rule, size=5)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    text = footer.add_run("MAGHRABI  |  " + ("صفحة " if rtl else "Page "))
    text.font.size = Pt(8)
    text.font.color.rgb = _color(theme.muted)
    _field(footer, "PAGE")
    _paragraph_direction(footer, rtl)


def _sections(doc, content: Content, theme: Theme, rtl: bool):
    for number, section in enumerate(content.sections, 1):
        kicker = doc.add_paragraph()
        kicker.paragraph_format.space_before = Pt(13)
        kicker.paragraph_format.space_after = Pt(2)
        kicker.paragraph_format.keep_with_next = True
        label = kicker.add_run(f"{number:02d}  /  " + ("المحور" if rtl else "SECTION"))
        label.font.size = Pt(9)
        label.font.bold = True
        label.font.color.rgb = _color(theme.accent)
        _paragraph_direction(kicker, rtl)

        heading = doc.add_heading(section.title, level=1)
        _bookmark(heading, number)
        _paragraph_direction(heading, rtl)
        _rule(heading, theme.accent, size=12)

        for part in section.body:
            paragraph = doc.add_paragraph(part)
            paragraph.paragraph_format.space_after = Pt(9)
            paragraph.paragraph_format.line_spacing = 1.48
            paragraph.paragraph_format.widow_control = True
            _paragraph_direction(paragraph, rtl)


def _template_overview(doc, content: Content, theme: Theme, rtl: bool, document_type: str):
    """Visual index of user-supplied headings only; does not fabricate project data."""
    if document_type not in ("study", "strategy"):
        return
    heading = {
        "study": ("خارطة محاور الدراسة", "STUDY ROADMAP"),
        "strategy": ("محاور الخطة الاستراتيجية", "STRATEGIC PILLARS"),
    }[document_type][0 if rtl else 1]
    p = doc.add_heading(heading, 1)
    _paragraph_direction(p, rtl)
    _rule(p, theme.accent, size=11)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(.72)
    table.columns[1].width = Inches(5.70)
    for i, item in enumerate(content.sections, 1):
        cells = table.add_row().cells
        cells[0].text = f"{i:02d}"
        cells[1].text = item.title
        for j, cell in enumerate(cells):
            _set_cell_margins(cell, top=110, bottom=110, start=160, end=160)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                _shade_cell(cell, theme.pale)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                _paragraph_direction(paragraph, rtl)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = _color(theme.accent if j == 0 else theme.dark)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def _formal_letter(doc, content: Content, theme: Theme, rtl: bool):
    """Compact correspondence layout; never invent recipient, date or signature."""
    section = doc.sections[0]
    # Letterhead is deliberately unbranded; no fictitious government letterhead.
    section.top_margin = Inches(.92)
    section.bottom_margin = Inches(.86)
    section.different_first_page_header_footer = False
    header = section.header.paragraphs[0]
    _rule(header, theme.accent, size=14)
    _paragraph_direction(header, rtl)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(20)
    title.paragraph_format.keep_with_next = True
    subject = title.add_run(
        ("الموضوع: " if rtl else "Subject: ") + content.title
    )
    subject.bold = True
    subject.font.size = Pt(15)
    subject.font.color.rgb = _color(theme.dark)
    _paragraph_direction(title, rtl)
    _rule(title, theme.rule, size=7)
    if content.summary:
        opening = doc.add_paragraph(content.summary)
        opening.paragraph_format.space_after = Pt(17)
        _paragraph_direction(opening, rtl)

    for item in content.sections:
        lead = doc.add_paragraph()
        lead.paragraph_format.space_before = Pt(9)
        lead.paragraph_format.space_after = Pt(6)
        lead.paragraph_format.keep_with_next = True
        run = lead.add_run(item.title)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = _color(theme.secondary)
        _paragraph_direction(lead, rtl)
        for body in item.body:
            paragraph = doc.add_paragraph(body)
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.line_spacing = 1.52
            _paragraph_direction(paragraph, rtl)

    # A small footer uses neutral document page numbering, not an official seal.
    footer = section.footer.paragraphs[0]
    _paragraph_direction(footer, rtl)
    num = footer.add_run(("صفحة " if rtl else "Page "))
    num.font.size = Pt(8)
    num.font.color.rgb = _color(theme.muted)
    _field(footer, "PAGE")


def export_premium_docx(content: Content, design_style: str = "luxury",
                        document_type: str = "general") -> bytes:
    if design_style not in STYLE_NAMES:
        raise ValueError("Unsupported Word design style")
    if document_type not in DOCUMENT_TYPES:
        raise ValueError("Unsupported Word document type")
    theme = THEMES[design_style]
    rtl = content.language == "ar"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.96)
    section.bottom_margin = Inches(.86)
    section.left_margin = Inches(.90)
    section.right_margin = Inches(.90)
    section.header_distance = Inches(.38)
    section.footer_distance = Inches(.40)

    _setup_styles(doc, theme, rtl)
    if document_type == "letter":
        _formal_letter(doc, content, theme, rtl)
    else:
        _configure_header_footer(doc, theme, rtl)
        _cover(doc, content, theme, rtl, document_type)
        _contents(doc, content, theme, rtl)
        _summary(doc, content, theme, rtl, document_type)
        _template_overview(doc, content, theme, rtl, document_type)
        _sections(doc, content, theme, rtl)

    doc.core_properties.title = content.title
    doc.core_properties.subject = DOCUMENT_LABELS[document_type][0 if rtl else 1]
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

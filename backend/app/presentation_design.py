"""Premium, fully editable 16:9 PowerPoint layouts with safe Arabic RTL text."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import math
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .content import Content, Section

WIDTH, HEIGHT = 13.333, 7.5
STYLE_NAMES = ("luxury", "corporate", "minimal")


@dataclass(frozen=True)
class Palette:
    dark: str
    accent: str
    accent2: str
    paper: str
    white: str
    ink: str
    secondary: str
    line: str
    tint: str


THEMES = {
    "luxury": Palette("17212F", "C6A46A", "8BAAB0", "F9F7F2", "FFFFFF",
                      "1A2839", "65727E", "DED8CC", "F1EDE4"),
    "corporate": Palette("10364B", "11979D", "83C8CE", "F4F8FA", "FFFFFF",
                         "17374A", "5C7989", "D7E5E8", "E7F2F4"),
    "minimal": Palette("26273D", "8474CB", "D6C9F1", "F8F7FC", "FFFFFF",
                       "24273D", "6F7086", "E4E0F0", "F0EDFA"),
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _rect(slide, x: float, y: float, w: float, h: float, fill: str,
          *, radius: bool = False, stroke: str | None = None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if stroke:
        shape.line.color.rgb = _rgb(stroke)
        shape.line.width = Pt(.9)
    else:
        shape.line.fill.background()
    if radius:
        # Shape stays editable. Small corner radius is suitable for editorial cards.
        try:
            shape.adjustments[0] = 0.10
        except (AttributeError, IndexError, ValueError):
            pass
    return shape


def _ring(slide, x: float, y: float, diameter: float, color: str, width: float = 1.4):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                   Inches(diameter), Inches(diameter))
    shape.fill.background()
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width)


def _text(slide, value: str, x: float, y: float, w: float, h: float,
          size: float, color: str, rtl: bool, *, bold: bool = False,
          alignment: PP_ALIGN | None = None, vertical=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.025)
    frame.margin_top = frame.margin_bottom = Inches(.025)
    frame.vertical_anchor = vertical
    # Keep user text as a native PowerPoint text run, not a rasterized picture.
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = alignment or (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)
    paragraph._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.18
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
    return box


def _new(deck: Presentation, bg: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(bg)
    return slide


def _short_label(rtl: bool, key: str):
    words = {
        "contents": ("محاور العرض", "CONTENTS"),
        "summary": ("لمحة عن العرض", "OVERVIEW"),
        "continued": ("— تابع", "— continued"),
        "chapter": ("المحور", "CHAPTER"),
        "page": ("صفحة", "PAGE"),
    }
    return words[key][0 if rtl else 1]


def _brand(slide, theme: Palette, rtl: bool, *, dark: bool = False):
    fg = theme.accent2 if dark else theme.secondary
    _text(slide, "MAGHRABI  /  OFFICE STUDIO", .76, 7.06, 5.0, .22, 9, fg, False)


def _foot(slide, theme: Palette, rtl: bool, number: int):
    _rect(slide, .77, 6.87, 11.8, .012, theme.line)
    _brand(slide, theme, rtl)
    _text(slide, f"{number:02d}", 11.85, 7.015, .72, .26,
          11, theme.accent, False, bold=True, alignment=PP_ALIGN.RIGHT)


def _cover(deck: Presentation, content: Content, theme: Palette, rtl: bool):
    slide = _new(deck, theme.dark)
    _rect(slide, 0, 0, WIDTH, .10, theme.accent)
    _rect(slide, .82, .83, 1.20, .045, theme.accent)
    _ring(slide, 0.9 if rtl else 9.9, 1.3, 2.8, theme.accent, 2.2)
    _ring(slide, 1.42 if rtl else 10.42, 1.82, 1.75, theme.accent2, 1.0)
    _rect(slide, 2.06 if rtl else 9.75, 3.7, .11, 1.30, theme.accent)
    _text(slide, "MAGHRABI / PREMIUM PRESENTATION", 1.0, 1.08,
          11.35, .35, 12, theme.accent, rtl, bold=True)
    title_size = 40 if len(content.title) <= 52 else 33 if len(content.title) <= 105 else 27
    _text(slide, content.title, 3.15 if rtl else 1.0, 2.0,
          9.05, 2.22, title_size, theme.white, rtl, bold=True)
    if content.summary:
        # A long summary is fully reproduced on a dedicated editorial page.
        summary = content.summary if len(content.summary) <= 220 else _first_words(content.summary, 190) + "…"
        _text(slide, summary, 3.15 if rtl else 1.0, 4.83,
              8.95, 1.18, 17 if len(summary) <= 180 else 15,
              "DFE5E9", rtl)
    _rect(slide, 1.0, 6.67, 11.25, .018, "48606B")
    _brand(slide, theme, rtl, dark=True)
    return slide


def _first_words(text: str, maximum: int) -> str:
    if len(text) <= maximum:
        return text
    candidate = text[:maximum].rsplit(" ", 1)[0]
    return candidate or text[:maximum]


def _header(slide, heading: str, theme: Palette, rtl: bool,
            eyebrow: str, *, number: int | None = None):
    _rect(slide, 0, 0, WIDTH, .075, theme.accent)
    side = 12.38 if rtl else .72
    _rect(slide, side, .66, .09, .88, theme.accent)
    _text(slide, eyebrow, .92, .62, 11.5, .28, 10.5, theme.accent,
          rtl, bold=True)
    size = 29 if len(heading) < 53 else 25 if len(heading) < 100 else 21
    _text(slide, heading, .92, 1.02, 11.45, 1.17,
          size, theme.ink, rtl, bold=True)


def _overview(deck: Presentation, content: Content, theme: Palette, rtl: bool):
    if not content.summary or len(content.summary) <= 220:
        return
    # Paginate long summaries; never squeeze 600 characters into a small box.
    for offset, summary_part in enumerate(_split_text(content.summary, 295)):
        slide = _new(deck, theme.paper)
        heading = _short_label(rtl, "summary")
        if offset:
            heading += f" {_short_label(rtl, 'continued')}"
        _header(slide, heading, theme, rtl, "EXECUTIVE OVERVIEW")
        _rect(slide, .89, 2.55, 11.52, 3.74, theme.white, radius=True, stroke=theme.line)
        _rect(slide, 12.13 if rtl else 1.06, 2.94, .07, 2.90, theme.accent)
        _text(slide, summary_part, 1.38, 2.86, 10.47, 3.12, 19,
              theme.ink, rtl)
        _foot(slide, theme, rtl, len(deck.slides))


def _contents(deck: Presentation, content: Content, theme: Palette, rtl: bool):
    # No invented information: only approved section titles and their indices.
    sections = list(enumerate(content.sections, 1))
    per_page = 8
    for start in range(0, len(sections), per_page):
        slide = _new(deck, theme.paper)
        _header(slide, _short_label(rtl, "contents"), theme, rtl,
                "PRESENTATION ROADMAP")
        for slot, (number, section) in enumerate(sections[start:start + per_page]):
            col, row = slot % 2, slot // 2
            x = .88 + col * 6.0
            y = 2.40 + row * 1.03
            _rect(slide, x, y, 5.65, .85, theme.white, radius=True,
                  stroke=theme.line)
            badge_x = x + 4.91 if rtl else x + .20
            _rect(slide, badge_x, y + .22, .52, .42, theme.tint, radius=True)
            _text(slide, f"{number:02d}", badge_x + .045, y + .29, .44, .26,
                  11.5, theme.accent, False, bold=True,
                  alignment=PP_ALIGN.CENTER)
            title_x = x + .27 if rtl else x + .93
            _text(slide, section.title, title_x, y + .13, 4.55, .67,
                  14 if len(section.title) <= 58 else 12 if len(section.title) <= 99 else 10,
                  theme.ink, rtl, bold=True)
        _foot(slide, theme, rtl, len(deck.slides))


def _split_text(value: str, limit: int = 325) -> list[str]:
    """Keep each original word in order, even for very long paragraphs."""
    if len(value) <= limit:
        return [value]
    tokens = re.findall(r"\S+\s*|\s+", value)
    result, bucket = [], ""
    for token in tokens:
        if len(token) > limit:
            if bucket:
                result.append(bucket.strip())
                bucket = ""
            while len(token) > limit:
                result.append(token[:limit])
                token = token[limit:]
        if bucket and len(bucket) + len(token) > limit:
            result.append(bucket.strip())
            bucket = ""
        bucket += token
    if bucket:
        result.append(bucket.strip())
    return [item for item in result if item]


def _layout_groups(section: Section) -> list[list[str]]:
    units = [part for body in section.body for part in _split_text(body)]
    pages: list[list[str]] = []
    current: list[str] = []
    units_used = 0
    for item in units:
        # One generous feature panel or two/three balanced cards.
        lines = max(1, math.ceil(len(item) / 48))
        max_items = 3 if len(item) <= 140 else 2
        if current and (units_used + lines > 9 or len(current) >= max_items):
            pages.append(current)
            current, units_used = [], 0
        current.append(item)
        units_used += lines
    if current:
        pages.append(current)
    return pages


def _card(slide, theme: Palette, rtl: bool, text: str, x: float, y: float,
          w: float, h: float, number: int, *, simple=False):
    _rect(slide, x, y, w, h, theme.white, radius=True, stroke=theme.line)
    accent_x = x + w - .16 if rtl else x + .08
    _rect(slide, accent_x, y + .28, .045, h - .57, theme.accent)
    index_x = x + w - .64 if rtl else x + .26
    _text(slide, f"{number:02d}", index_x, y + .33, .39, .42, 15,
          theme.accent, False, bold=True)
    text_y = y + (.82 if not simple else .68)
    budget = h - (text_y - y) - .21
    if w < 4.0:
        size = 18 if len(text) <= 75 else 15 if len(text) <= 145 else 13
    elif w < 6.4:
        size = 20 if len(text) <= 120 else 17 if len(text) <= 210 else 15
    else:
        size = 21 if len(text) <= 215 else 18 if len(text) <= 325 else 16
    _text(slide, text, x + .35, text_y, w - .73, budget, size, theme.ink, rtl)


def _content_slide(deck: Presentation, section: Section, section_index: int,
                   page_index: int, pieces: list[str], theme: Palette, rtl: bool):
    slide = _new(deck, theme.paper)
    subtitle = f"{_short_label(rtl, 'chapter')} {section_index:02d}  /  {len(deck.slides):02d}"
    title = section.title + (f" {_short_label(rtl, 'continued')}" if page_index else "")
    _header(slide, title, theme, rtl, subtitle)
    if len(pieces) == 1:
        # Editorial feature: oversized section number + readable copy panel.
        circle_x = .95 if rtl else 10.10
        _ring(slide, circle_x, 3.12, 2.20, theme.line, 3.0)
        _text(slide, f"{section_index:02d}", circle_x + .39, 3.66, 1.42, .94,
              48, theme.accent, False, bold=True, alignment=PP_ALIGN.CENTER)
        panel_x = 3.85 if rtl else .85
        _card(slide, theme, rtl, pieces[0], panel_x, 2.43, 8.51, 3.94, 1)
    elif len(pieces) == 2:
        for i, piece in enumerate(pieces):
            _card(slide, theme, rtl, piece,
                  .88 + i * 6.04, 2.46, 5.72, 3.91, i + 1)
    else:
        for i, piece in enumerate(pieces):
            _card(slide, theme, rtl, piece,
                  .86 + i * 4.16, 2.49, 3.91, 3.86, i + 1)
    # Explicit notes stay available for each continuation slide.
    if section.notes:
        slide.notes_slide.notes_text_frame.text = section.notes
    _foot(slide, theme, rtl, len(deck.slides))


def export_premium_pptx(content: Content, design_style: str = "luxury") -> bytes:
    if design_style not in STYLE_NAMES:
        raise ValueError("Unsupported PowerPoint design style")
    theme = THEMES[design_style]
    rtl = content.language == "ar"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    _cover(deck, content, theme, rtl)
    _overview(deck, content, theme, rtl)
    _contents(deck, content, theme, rtl)
    for section_idx, section in enumerate(content.sections, 1):
        for page_idx, group in enumerate(_layout_groups(section)):
            _content_slide(deck, section, section_idx, page_idx, group, theme, rtl)
    out = BytesIO()
    deck.save(out)
    return out.getvalue()

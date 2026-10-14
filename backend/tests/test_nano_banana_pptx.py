"""Opt-in visual PowerPoint mode must be authenticated and retain editable text."""
from __future__ import annotations

import base64
from io import BytesIO

import httpx
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app.content import Content
from app.main import app
from app import presentation_visuals


def _art() -> bytes:
    output = BytesIO()
    Image.new("RGB", (320, 180), "#21405a").save(output, format="JPEG")
    return output.getvalue()


def _content() -> Content:
    return Content(
        title="خطة تحسين الخدمات",
        summary="عرض رسمي لتحسين الخدمة",
        language="ar",
        sections=[
            {"title": "الخدمات", "body": ["تطوير الإجراءات وقياس النتائج الفعلية."], "notes": "ملاحظات خاصة لا تُرسل."}
        ],
    )


def _shapes_text(deck: Presentation) -> str:
    return " ".join(
        shape.text for slide in deck.slides for shape in slide.shapes
        if shape.has_text_frame
    )


def test_visual_mode_requires_ai_login_and_preserves_all_native_text(monkeypatch):
    monkeypatch.setenv("AI_PROVIDER", "gemini")
    monkeypatch.setenv("GEMINI_API_KEY", "fake-key")
    monkeypatch.setenv("GEMINI_MODEL", "fake-text-model")
    monkeypatch.setenv("STUDIO_ACCESS_TOKEN", "secure-test-value")
    calls = []

    def fake_assets(content, style):
        calls.append((content.title, style))
        return {"cover": _art(), "section": _art()}

    monkeypatch.setattr("app.main.generate_visual_assets", fake_assets)
    payload = _content().model_dump()
    with TestClient(app) as client:
        unauth = client.post("/api/v1/ai/exports/pptx", json=payload,
                             headers={"X-Maghrabi-Client": "studio"})
        assert unauth.status_code == 401
        assert not calls

        # Existing editable export stays free of image-generation calls.
        plain = client.post("/api/v1/exports/pptx", json=payload)
        assert plain.status_code == 200
        assert not calls
        assert not any(shape.shape_type == 13 for slide in Presentation(BytesIO(plain.content)).slides
                       for shape in slide.shapes)

        client.post("/api/v1/ai/login", json={"access_token": "secure-test-value"})
        assert client.post("/api/v1/ai/exports/pptx", json=payload).status_code == 401
        visual = client.post(
            "/api/v1/ai/exports/pptx?design_style=corporate",
            json=payload, headers={"X-Maghrabi-Client": "studio"},
        )
        assert visual.status_code == 200
        assert calls == [(payload["title"], "corporate")]
        deck = Presentation(BytesIO(visual.content))
        assert "خطة تحسين الخدمات" in _shapes_text(deck)
        assert "تطوير الإجراءات وقياس النتائج الفعلية." in _shapes_text(deck)
        assert deck.slides[-1].notes_slide.notes_text_frame.text == "ملاحظات خاصة لا تُرسل."
        pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) >= 3  # cover, soft content backdrop, and inline editorial artwork
        assert all(pic.has_text_frame is False for pic in pictures)


def test_gemini_image_request_contains_no_notes_or_body(monkeypatch):
    content = _content()
    prompt = presentation_visuals._prompt(content, "luxury", cover=False)
    assert "الخدمات" in prompt
    assert "تطوير الإجراءات" not in prompt
    assert "ملاحظات خاصة" not in prompt

    seen = []
    encoded = base64.b64encode(_art()).decode()
    def respond(request: httpx.Request):
        seen.append(request)
        return httpx.Response(
            200, json={"candidates": [{"content": {"parts": [
                {"inlineData": {"mimeType": "image/jpeg", "data": encoded}}
            ]}}]},
        )

    with httpx.Client(transport=httpx.MockTransport(respond)) as client:
        result = presentation_visuals._generate(client, "gemini-3-pro-image", "fake-key", prompt)
    assert result[:2] == bytes((0xFF, 0xD8))
    assert Image.open(BytesIO(result)).size == (1920, 1080)
    assert len(seen) == 1
    assert seen[0].headers["x-goog-api-key"] == "fake-key"
    body = seen[0].content.decode()
    assert "2K" in body and "16:9" in body
    assert "تطوير الإجراءات" not in body


def test_missing_image_key_fails_without_paid_request(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    from fastapi import HTTPException
    try:
        presentation_visuals.generate_visual_assets(_content(), "luxury")
        assert False, "must reject absent image API key"
    except HTTPException as exc:
        assert exc.status_code == 503

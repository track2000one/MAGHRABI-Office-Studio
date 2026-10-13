"""Branded Word documents retain editable content and securely validate uploaded logos."""
from io import BytesIO
from zipfile import ZipFile
import json

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.main import app


CONTENT = {
    "title": "دراسة تطوير الخدمات",
    "summary": "الملخص المعتمد من المستخدم.",
    "language": "ar",
    "sections": [
        {"title": "الوضع الراهن", "body": ["نتائج أولية واردة من المستخدم."], "notes": ""},
        {"title": "التوصيات", "body": ["نوصي بمراجعة الإجراءات."], "notes": ""},
    ],
}
BRAND = {
    "organization_name": "جهة المستخدم للاختبار",
    "header_text": "الإدارة الفنية",
    "footer_text": "نسخة داخلية للاختبار",
    "table_title": "جدول مهام العمل",
    "table_rows": [
        ["المهمة", "المسؤول", "الحالة"],
        ["إعداد التقرير", "الموظف المكلف", "قيد الدراسة"],
        ["مراجعة النتائج", "المراجعة", "لم تبدأ"],
    ],
}


def png_logo():
    out = BytesIO()
    Image.new("RGBA", (200, 110), (10, 90, 110, 255)).save(out, format="PNG")
    return out.getvalue()


def custom_request(client, customization=None, file_bytes=None, *, kind="report", style="corporate"):
    files = {}
    if file_bytes is not None:
        files["logo"] = ("agency.png", file_bytes, "image/png")
    return client.post(
        f"/api/v1/exports/docx/custom?design_style={style}&document_type={kind}",
        data={
            "content_json": json.dumps(CONTENT, ensure_ascii=False),
            "customization_json": json.dumps(customization if customization is not None else BRAND, ensure_ascii=False),
        },
        files=files or None,
    )


def test_word_custom_brand_logo_and_tabular_data_are_native_and_editable():
    with TestClient(app) as client:
        response = custom_request(client, file_bytes=png_logo())
    assert response.status_code == 200, response.text[:800]
    document = Document(BytesIO(response.content))
    paragraphs = [p.text for p in document.paragraphs]
    assert any(p.strip() == "جهة المستخدم للاختبار" for p in paragraphs)
    assert "جدول مهام العمل" in paragraphs
    assert CONTENT["sections"][0]["body"][0] in paragraphs
    assert document.sections[0].header.paragraphs[0].text.find("الإدارة الفنية") >= 0
    assert "نسخة داخلية للاختبار" in document.sections[0].footer.paragraphs[0].text
    assert any(table.cell(1, 0).text == "إعداد التقرير" for table in document.tables if len(table.rows) > 1)
    with ZipFile(BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert any(name.startswith("word/media/") and name.endswith(".png") for name in names)
        xml = archive.read("word/document.xml")
        assert b"w:tblHeader" in xml
        assert b"w:bidi" in xml


def test_custom_letter_branding_does_not_add_cover_or_unprovided_official_data():
    with TestClient(app) as client:
        response = custom_request(client, customization={
            "organization_name": "مختبر التطوير",
            "footer_text": "سري للاستخدام الداخلي",
        }, file_bytes=png_logo(), kind="letter", style="luxury")
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "الموضوع: دراسة تطوير الخدمات" in text
    assert "المحتويات" not in text
    assert "MAGHRABI" not in text
    assert "سري للاستخدام الداخلي" in document.sections[0].footer.paragraphs[0].text
    assert "مختبر التطوير" in document.sections[0].header.paragraphs[0].text
    assert not document.tables


def test_reject_wrong_file_type_even_if_marked_image_png():
    with TestClient(app) as client:
        response = custom_request(client, file_bytes=b"%PDF-1.7 invalid image")
    assert response.status_code == 422
    assert "PNG" in response.text or "JPEG" in response.text


def test_reject_oversized_logo_and_uneven_table_columns():
    with TestClient(app) as client:
        big = custom_request(client, file_bytes=b"a" * (1024 * 1024 + 1))
        assert big.status_code == 422
        uneven = custom_request(client, customization={
            "table_title": "الجدول",
            "table_rows": [["أ", "ب"], ["فقط"]],
        })
        assert uneven.status_code == 422
        with_bad_column = custom_request(client, customization={
            "table_rows": [["واحد", "اثنان", "ثلاثة", "أربعة", "خمسة", "ستة", "سبعة"],
                           ["1", "2", "3", "4", "5", "6", "7"]],
        })
        assert with_bad_column.status_code == 422


def test_old_json_docx_export_remains_compatible():
    with TestClient(app) as client:
        response = client.post("/api/v1/exports/docx?document_type=report", json=CONTENT)
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    assert CONTENT["title"] in [p.text for p in doc.paragraphs]

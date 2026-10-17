"""Lecture quiz security, text extraction, grounded MCQs and API regression tests."""
from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
import fitz
import pytest

from app import quiz
from app.main import app


SOURCE = (
    "The cardiovascular system contains the heart, blood vessels, and blood. "
    "The heart sends blood through arteries to tissues. Veins return blood to the heart. "
    "Capillaries are the smallest vessels and are the site of gas and nutrient exchange. "
    "Oxygen binds to hemoglobin inside red blood cells. "
    "The pulmonary circulation sends blood from the heart to the lungs and back. "
) * 3
QUOTES = [
    "The heart sends blood through arteries to tissues.",
    "Veins return blood to the heart.",
    "Capillaries are the smallest vessels",
    "Oxygen binds to hemoglobin inside red blood cells.",
    "The pulmonary circulation sends blood from the heart to the lungs and back.",
]


def docx_data():
    document = Document()
    document.add_heading("Lecture on circulation", 0)
    document.add_paragraph(SOURCE)
    document.add_paragraph("Additional study note: " + SOURCE)
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def pdf_data():
    doc = fitz.open()
    p = doc.new_page()
    p.insert_textbox(fitz.Rect(50, 50, 550, 780), SOURCE[:600], fontsize=10)
    p2 = doc.new_page()
    p2.insert_textbox(fitz.Rect(50, 50, 550, 780), SOURCE[600:], fontsize=10)
    result = doc.tobytes()
    doc.close()
    return result


def mock_bank(count=5, quote=None):
    questions=[]
    for i in range(count):
        questions.append({
            "question": f"What is the primary fact number {i+1} supported by the lecture?",
            "options": ["Option Alpha","Option Beta","Option Gamma","Option Delta"],
            "correct_index": i % 4, "explanation": "This is the correct choice because the lecture supports it.",
            "source_quote": quote or QUOTES[i % len(QUOTES)], "source_page": 1,
        })
    return {"title": "Practice from circulation lecture", "questions": questions}


def test_extract_docx_retains_paragraphs_and_page_context():
    out = quiz.extract_lecture(docx_data(), "Circulation.docx")
    assert out.pages == 1
    assert "[PAGE 1]" in out.source_text
    assert "Capillaries" in out.source_text
    assert out.truncated is False


def test_extract_pdf_retains_page_markers():
    out=quiz.extract_lecture(pdf_data(), "Circulation.pdf")
    assert out.pages==2
    assert "[PAGE 1]" in out.source_text
    assert "[PAGE 2]" in out.source_text


def test_invalid_and_image_only_files_fail_without_ocr():
    with pytest.raises(Exception) as error:
        quiz.extract_lecture(b"not a real pdf", "report.pdf")
    assert error.value.status_code==422
    with pytest.raises(Exception) as invalid:
        quiz.extract_lecture(b"12345678", "report.doc")
    assert invalid.value.status_code==415
    scanned=fitz.open()
    scanned.new_page()
    raw=scanned.tobytes()
    scanned.close()
    with pytest.raises(Exception) as blank:
        quiz.extract_lecture(raw,"scanned.pdf")
    assert blank.value.status_code==422
    assert "قابل" in str(blank.value.detail)


def test_quiz_rejects_wrong_count_and_duplicate_options_but_not_unmatched_evidence(monkeypatch):
    monkeypatch.setattr(quiz.ai,"status",lambda:{"configured":True})
    monkeypatch.setattr(quiz, "_provider_quiz",lambda client,payload:json.dumps(mock_bank(5)))
    request=quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=10)
    with pytest.raises(Exception) as count_error:
        quiz.create_quiz(request)
    assert count_error.value.status_code==502
    monkeypatch.setattr(quiz,"_provider_quiz",lambda client,payload:json.dumps(mock_bank(5,quote="Quote never present in user lecture")))
    created=quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=5))
    assert len(created.questions)==5
    assert all(q.source_verified is False for q in created.questions)
    assert all(q.source_quote == "" and q.source_page == 0 for q in created.questions)
    bad=mock_bank(5)
    bad["questions"][0]["options"]=["same","same","different","another"]
    monkeypatch.setattr(quiz,"_provider_quiz",lambda client,payload:json.dumps(bad))
    with pytest.raises(Exception) as options_error:
        quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=5))
    assert options_error.value.status_code==502


def test_question_evidence_page_is_derived_not_trusted(monkeypatch):
    monkeypatch.setattr(quiz.ai,"status",lambda:{"configured":True})
    mocked=mock_bank(5)
    mocked["questions"][0]["source_page"]=80
    monkeypatch.setattr(quiz,"_provider_quiz",lambda client,payload:json.dumps(mocked))
    data=quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=5))
    assert len(data.questions)==5
    assert data.questions[0].source_page==1
    assert data.questions[0].source_verified is True


def test_endpoint_requires_ai_session_and_reuses_upload_text(monkeypatch):
    monkeypatch.setenv("STUDIO_ACCESS_TOKEN","strong-test-office-token")
    monkeypatch.setenv("AI_PROVIDER","gemini")
    monkeypatch.setenv("GEMINI_API_KEY","placeholder-only")
    monkeypatch.setenv("GEMINI_MODEL","placeholder")
    monkeypatch.setattr(quiz,"_provider_quiz",lambda client,payload:json.dumps(mock_bank(payload["count"])))
    with TestClient(app) as client:
        unauth=client.post("/api/v1/ai/quiz/extract",files={"file":("lecture.docx",docx_data(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert unauth.status_code==401
        login=client.post("/api/v1/ai/login",json={"access_token":"strong-test-office-token"})
        assert login.status_code==200
        extracted=client.post("/api/v1/ai/quiz/extract",headers={"X-Maghrabi-Client":"studio"},
                      files={"file":("lecture.docx",docx_data(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert extracted.status_code==200,extracted.text
        text=extracted.json()["source_text"]
        generated=client.post("/api/v1/ai/quiz/generate",headers={"X-Maghrabi-Client":"studio"},
                         json={"source_text":text,"count":5,"difficulty":"hard","language":"en"})
        assert generated.status_code==200,generated.text
        assert len(generated.json()["questions"])==5
        assert generated.json()["questions"][0]["source_quote"] in text


@pytest.mark.parametrize("count,difficulty",[(0,"easy"),(100,"easy"),(5,"invalid"),(2,"mixed")])
def test_invalid_quiz_settings_are_rejected(count,difficulty):
    with pytest.raises(ValueError):
        quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=count,difficulty=difficulty)


def test_pdf_typography_and_line_wraps_are_supported_without_paraphrasing():
    pages = {
        7: "The findings of the ﬁnal workplace-research report were recorded "
           "in a written report. Employees’ responses were analyzed.",
        13: "This unrelated page discusses open public materials.",
    }
    hit = quiz._find_evidence(
        "the findings of the final workplace research report were recorded "
        "in a written report",
        pages,
    )
    assert hit is not None
    assert hit[0] == 7
    assert "ﬁnal" in hit[1]
    assert "workplace-research" in hit[1]
    assert quiz._find_evidence(
        "The findings were published in an oral presentation.", pages,
    ) is None
    assert quiz._find_evidence("in a made up document", pages) is None


def test_unicode_case_spacing_and_arabic_diacritics_resolve_original_evidence():
    pages = {
        2: "The organization’s FINAL decision was written in the report.",
        4: "تُكتَبُ النَّتائِجُ في التَّقريرِ الرَّسمِي للمؤسسة.",
    }
    assert quiz._find_evidence("The organization's final decision was written in the report", pages)[0] == 2
    match = quiz._find_evidence("تكتب النتائج في التقرير الرسمي", pages)
    assert match is not None
    assert match[0] == 4
    assert "التَّقريرِ" in match[1]


def test_unverifiable_quote_does_not_retry_or_block_valid_exam(monkeypatch):
    monkeypatch.setattr(quiz.ai, "status", lambda: {"configured": True})
    calls = []
    def provider(client, payload):
        calls.append(payload.copy())
        return json.dumps(mock_bank(5, quote="Unsupported medical claims not in the lecture."))
    monkeypatch.setattr(quiz, "_provider_quiz", provider)
    generated = quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 1]\n" + SOURCE, count=5))
    assert len(generated.questions) == 5
    assert len(calls) == 1
    assert all(q.source_quote == "" and q.source_page == 0 and not q.source_verified
               for q in generated.questions)


def test_mixed_verified_and_unverified_evidence_are_labelled_individually(monkeypatch):
    monkeypatch.setattr(quiz.ai, "status", lambda: {"configured": True})
    data = mock_bank(5)
    data["questions"][0]["source_quote"] = "A fabricated quote that is absent from the uploaded PDF"
    data["questions"][1]["source_quote"] = ""
    monkeypatch.setattr(quiz, "_provider_quiz", lambda client, payload: json.dumps(data))
    res = quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 2]\n" + SOURCE, count=5))
    assert len(res.questions) == 5
    assert not res.questions[0].source_verified and res.questions[0].source_page == 0
    assert not res.questions[1].source_verified and res.questions[1].source_quote == ""
    assert res.questions[2].source_verified
    assert res.questions[2].source_quote in SOURCE
    assert res.questions[2].source_page == 2


def test_invalid_options_still_fail_after_one_structure_retry(monkeypatch):
    monkeypatch.setattr(quiz.ai, "status", lambda: {"configured": True})
    malformed = mock_bank(5)
    malformed["questions"][0]["options"] = ["identical", "identical", "different", "last"]
    calls=[]
    def provider(client, payload):
        calls.append(payload)
        return json.dumps(malformed)
    monkeypatch.setattr(quiz, "_provider_quiz", provider)
    with pytest.raises(Exception) as error:
        quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 1]\n" + SOURCE, count=5))
    assert error.value.status_code==502
    assert len(calls)==2


def test_page_reference_is_canonicalized_using_actual_pdf_page_number(monkeypatch):
    monkeypatch.setattr(quiz.ai, "status", lambda: {"configured": True})
    data = mock_bank(5)
    data["questions"][0]["source_quote"] = "THE HEART SENDS BLOOD THROUGH ARTERIES TO TISSUES"
    data["questions"][0]["source_page"] = 22
    monkeypatch.setattr(quiz, "_provider_quiz", lambda client, payload: json.dumps(data))
    res = quiz.create_quiz(quiz.QuizRequest(source_text="[PAGE 6]\n" + SOURCE, count=5))
    assert res.questions[0].source_page == 6
    assert res.questions[0].source_quote == "The heart sends blood through arteries to tissues"


def test_generate_40_question_exam_as_four_verified_ten_question_batches(monkeypatch):
    """Each independent request stays small, carries prior questions, and spans 40."""
    monkeypatch.setattr(quiz.ai,"status",lambda:{"configured":True})
    seen_request_lengths = []
    asked = []
    def provider(client,payload):
        batch_size = payload["count"]
        previous = payload.get("previous_questions",[])
        seen_request_lengths.append((batch_size,len(previous)))
        start = len(previous)
        questions=[]
        for offset in range(batch_size):
            idx = start + offset + 1
            questions.append({
                "question":f"Which circulation fact describes item number {idx} from the lecture?",
                "options":["A different item","The item in the lecture","Another choice","Unrelated answer"],
                "correct_index":1,
                "explanation":"The cardiovascular lecture explicitly describes the relevant system.",
                "source_quote":QUOTES[(idx - 1) % len(QUOTES)],
                "source_page":1,
            })
        asked.extend([question["question"] for question in questions])
        return json.dumps({"title":"Circulatory system practice test","questions":questions})

    monkeypatch.setattr(quiz,"_provider_quiz",provider)
    built=[]
    for size in (10,10,10,10):
        req=quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=size,avoid_questions=[q.question for q in built])
        batch=quiz.create_quiz(req)
        built.extend(batch.questions)
    assert len(built)==40
    assert len({q.question for q in built})==40
    assert seen_request_lengths==[(10,0),(10,10),(10,20),(10,30)]
    assert all(q.source_quote in SOURCE for q in built)


def test_large_exam_batch_rejects_repeated_question_from_previous_batch(monkeypatch):
    monkeypatch.setattr(quiz.ai,"status",lambda:{"configured":True})
    generated = mock_bank(5)
    generated["questions"][0]["question"] = "  " + generated["questions"][0]["question"].upper() + "!!"
    seen=[]
    def provider(client,payload):
        seen.append(payload)
        return json.dumps(generated)
    monkeypatch.setattr(quiz,"_provider_quiz",provider)
    with pytest.raises(Exception) as error:
        quiz.create_quiz(quiz.QuizRequest(
            source_text="[PAGE 1]\n"+SOURCE,count=5,
            avoid_questions=[mock_bank(5)["questions"][0]["question"]],
        ))
    assert error.value.status_code==502
    assert len(seen)==2  # Existing one-retry strategy also applies to duplicates.
    assert "previous_questions" in seen[1]
    assert "validation_feedback" in seen[1]


def test_large_exam_batch_rejects_above_40_questions():
    with pytest.raises(ValueError):
        quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=10,
                         avoid_questions=[f"Previous question number {n}?" for n in range(31)])
    with pytest.raises(ValueError):
        quiz.QuizRequest(source_text="[PAGE 1]\n"+SOURCE,count=10,
                         avoid_questions=["a"] * 41)


def test_gemini_transient_503_retries_with_bounded_backoff(monkeypatch):
    import httpx
    monkeypatch.setenv("AI_PROVIDER", "gemini")
    monkeypatch.setenv("GEMINI_API_KEY", "dummy-test-key")
    monkeypatch.setenv("GEMINI_MODEL", "test-gemini-model")
    delays = []
    monkeypatch.setattr(quiz.time,"sleep",lambda seconds: delays.append(seconds))
    statuses = iter((503,503,200))
    attempts=[]
    def handler(request):
        assert "x-goog-api-key" in request.headers
        attempts.append(request)
        code=next(statuses)
        if code==200:
            return httpx.Response(200,json={"candidates":[{
                "finishReason":"STOP","content":{"parts":[{"text":json.dumps(mock_bank(5))}]}
            }]})
        return httpx.Response(code,json={"error":{"status":"UNAVAILABLE"}})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result=quiz._provider_quiz(client,{"count":5,"lecture_text":"test"})
    assert len(attempts)==3
    assert delays==[1.5,3.0]
    assert json.loads(result)["title"]=="Practice from circulation lecture"


def test_gemini_retry_after_is_capped_and_quota_429_never_retried(monkeypatch):
    import httpx
    monkeypatch.setenv("AI_PROVIDER","gemini")
    monkeypatch.setenv("GEMINI_API_KEY","dummy")
    monkeypatch.setenv("GEMINI_MODEL","test-model")
    delays=[]
    monkeypatch.setattr(quiz.time,"sleep",lambda value:delays.append(value))
    attempts=[]
    def busy(request):
        attempts.append(request)
        return httpx.Response(503,headers={"Retry-After":"120"},
                              json={"error":{"status":"UNAVAILABLE"}})
    with httpx.Client(transport=httpx.MockTransport(busy)) as client:
        with pytest.raises(Exception) as exc:
            quiz._provider_quiz(client,{"count":5,"lecture_text":"test"})
    assert exc.value.status_code==502
    assert "503" in exc.value.detail
    assert len(attempts)==3 and delays==[6.0,6.0]

    attempts.clear()
    delays.clear()
    def quota(request):
        attempts.append(request)
        return httpx.Response(429,json={"error":{"status":"RESOURCE_EXHAUSTED"}})
    with httpx.Client(transport=httpx.MockTransport(quota)) as client:
        with pytest.raises(Exception) as exc:
            quiz._provider_quiz(client,{"count":5,"lecture_text":"test"})
    assert exc.value.status_code==429
    assert len(attempts)==1
    assert delays==[]


def test_gemini_400_does_not_retry_transiently_but_can_fallback_schema(monkeypatch):
    import httpx
    monkeypatch.setenv("AI_PROVIDER","gemini")
    monkeypatch.setenv("GEMINI_API_KEY","dummy")
    monkeypatch.setenv("GEMINI_MODEL","test-model")
    delays=[]
    monkeypatch.setattr(quiz.time,"sleep",lambda value:delays.append(value))
    requests=[]
    def handler(request):
        requests.append(json.loads(request.content))
        if len(requests)==1:
            return httpx.Response(400,json={
                "error":{"message":"responseFormat is not supported"}})
        return httpx.Response(200,json={"candidates":[{
            "finishReason":"STOP","content":{"parts":[{"text":json.dumps(mock_bank(5))}]}
        }]})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result=quiz._provider_quiz(client,{"count":5,"lecture_text":"test"})
    assert len(requests)==2
    assert delays==[]
    assert "responseFormat" in requests[0]["generationConfig"]
    assert "responseMimeType" in requests[1]["generationConfig"]
    assert len(json.loads(result)["questions"])==5

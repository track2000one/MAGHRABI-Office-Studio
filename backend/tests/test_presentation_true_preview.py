"""True PPTX render and per-page artwork regression checks without external AI calls."""
from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import fitz
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app.content import Content
from app.exports import export_pptx
from app.main import app
from app.presentation_payload import PresentationArtwork, PresentationRenderRequest, compose_presentation


def _content():
    return Content(
        title="عرض رسمي تجريبي",
        summary="تحسين جودة الأعمال والإجراءات",
        language="ar",
        sections=[
            {"title": "الأهداف", "body": ["تعزيز الكفاءة التشغيلية."], "notes": "تعليق خاص"},
            {"title": "النتائج", "body": ["متابعة مؤشرات قابلة للقياس."], "notes": ""},
        ],
    )


def _art():
    buffer = BytesIO()
    Image.new("RGB", (640, 360), "#40557f").save(buffer, format="JPEG")
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()


def _all_text(slide):
    return [shape.text for shape in slide.shapes if shape.has_text_frame]


def _image_count(slide):
    return sum(shape.shape_type == 13 for shape in slide.shapes)


def test_exact_slide_override_keeps_all_other_pages_and_editable_text():
    content = _content()
    baseline = Presentation(BytesIO(export_pptx(content)))
    target = len(baseline.slides) - 1
    request = PresentationRenderRequest(
        content=content, artwork=PresentationArtwork(slides={str(target): _art()})
    )
    modified = Presentation(BytesIO(compose_presentation(request, "luxury")))
    assert len(modified.slides) == len(baseline.slides)
    for index in range(len(baseline.slides)):
        assert _all_text(modified.slides[index]) == _all_text(baseline.slides[index])
        if index == target:
            assert _image_count(modified.slides[index]) == _image_count(baseline.slides[index]) + 1
        else:
            assert _image_count(modified.slides[index]) == _image_count(baseline.slides[index])
    assert modified.slides[target].notes_slide.notes_text_frame.text == ""


def test_preview_endpoint_returns_real_pdf_pages_with_same_pptx_payload(monkeypatch):
    from app import presentation_preview
    original = _content()
    count = len(Presentation(BytesIO(export_pptx(original))).slides)
    invoked = []

    def fake_convert(command, **kwargs):
        invoked.append(command)
        assert "pdf:impress_pdf_Export" in command
        output = Path(command[-1]).with_suffix(".pdf")
        pdf = fitz.open()
        for i in range(count):
            page = pdf.new_page(width=960, height=540)
            page.insert_text((32, 40), f"Page {i + 1}")
        pdf.save(str(output))
        pdf.close()
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(presentation_preview.subprocess, "run", fake_convert)
    with TestClient(app) as client:
        payload = {"content": original.model_dump(), "artwork": {"slides": {}}}
        response = client.post("/api/v1/presentations/preview?design_style=corporate", json=payload)
        export = client.post("/api/v1/presentations/export?design_style=corporate", json=payload)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/pdf")
        assert response.headers["x-preview-pages"] == str(count)
        assert response.headers["cache-control"] == "no-store, private"
        assert export.status_code == 200
        assert len(Presentation(BytesIO(export.content)).slides) == count
        assert len(fitz.open(stream=response.content, filetype="pdf")) == count
        assert len(invoked) == 1


def test_invalid_images_and_indices_are_rejected():
    content = _content().model_dump()
    with TestClient(app) as client:
        malformed = client.post("/api/v1/presentations/export", json={
            "content": content, "artwork": {"cover": "data:image/jpeg;base64,not-valid!", "slides": {}},
        })
        assert malformed.status_code == 422
        stale_index = client.post("/api/v1/presentations/preview", json={
            "content": content, "artwork": {"slides": {"79": _art()}},
        })
        assert stale_index.status_code == 422
        invalid_key = client.post("/api/v1/presentations/export", json={
            "content": content, "artwork": {"slides": {"-1": _art()}},
        })
        assert invalid_key.status_code == 422


def test_redesign_endpoint_rejects_without_login_before_any_paid_call(monkeypatch):
    calls = []
    monkeypatch.setattr("app.main.generate_slide_artwork", lambda *args: calls.append(args))
    with TestClient(app) as client:
        result = client.post("/api/v1/ai/presentations/redesign-slide", json={
            "content": _content().model_dump(), "design_style": "luxury", "slide_index": 2,
        }, headers={"X-Maghrabi-Client": "studio"})
        assert result.status_code in (401, 503)
        assert not calls

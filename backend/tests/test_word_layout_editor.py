"""Document layout customization: alignment, dimensions, typography, safe enum validation."""
from io import BytesIO
import json
from zipfile import ZipFile

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from app.document_customization import WordCustomization
from app.main import app


CONTENT = {
    "title": "تقرير الأداء الشهري",
    "summary": "ملخص البيانات الموافق عليها من المستخدم.",
    "language": "ar",
    "sections": [
        {"title": "الأداء", "body": ["مؤشرات الأداء الموثقة."], "notes": ""},
        {"title": "التطوير", "body": ["خطة التحسين المقترحة."], "notes": ""},
    ],
}


def logo_bytes():
    out = BytesIO()
    Image.new("RGB", (240, 120), (80, 50, 110)).save(out, format="PNG")
    return out.getvalue()


def export(client, options, *, document_type="report", logo=False):
    data = {
        "content_json": json.dumps(CONTENT, ensure_ascii=False),
        "customization_json": json.dumps(options, ensure_ascii=False),
    }
    files = {"logo": ("test.png", logo_bytes(), "image/png")} if logo else None
    return client.post(
        f"/api/v1/exports/docx/custom?document_type={document_type}",
        data=data,
        files=files,
    )


def test_landscape_margins_type_scale_and_native_alignment():
    with TestClient(app) as client:
        response = export(client, {
            "organization_name": "مكتب الاختبار",
            "header_text": "الترويسة الأصلية",
            "footer_text": "تذييل الجهة",
            "logo_alignment": "center",
            "header_alignment": "center",
            "footer_alignment": "end",
            "page_orientation": "landscape",
            "page_margins": "compact",
            "body_font_size": 13,
            "show_page_numbers": False,
        }, logo=True)
    assert response.status_code == 200, response.text[:500]
    doc = Document(BytesIO(response.content))
    sec = doc.sections[0]
    assert sec.page_width > sec.page_height
    assert .65 < sec.left_margin.inches < .71
    assert doc.styles["Normal"].font.size.pt == 13
    assert sec.header.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert sec.footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert any(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in sec.header.paragraphs[1:])
    assert "تذييل الجهة" in sec.footer.paragraphs[0].text
    with ZipFile(BytesIO(response.content)) as archive:
        footer = next(archive.read(n) for n in archive.namelist() if n.startswith("word/footer") and n.endswith(".xml"))
        assert b"PAGE" not in footer
        assert b"w:jc" in footer


def test_logo_placement_cover_only_vs_header_only():
    with TestClient(app) as client:
        only_cover = export(client, {"logo_placement": "cover", "organization_name": "جهة الاختبار"}, logo=True)
        only_header = export(client, {"logo_placement": "header", "organization_name": "جهة الاختبار"}, logo=True)
    assert only_cover.status_code == 200
    assert only_header.status_code == 200
    cover_doc = Document(BytesIO(only_cover.content))
    header_doc = Document(BytesIO(only_header.content))
    assert any("graphic" in p._p.xml for p in cover_doc.paragraphs)
    assert not any("graphic" in p._p.xml for p in cover_doc.sections[0].header.paragraphs)
    assert not any("graphic" in p._p.xml for p in header_doc.paragraphs)
    assert any("graphic" in p._p.xml for p in header_doc.sections[0].header.paragraphs)


def test_spacious_portrait_and_letter_page_preferences():
    with TestClient(app) as client:
        response = export(client, {
            "organization_name": "إدارة المراسلات",
            "page_orientation": "portrait",
            "page_margins": "spacious",
            "footer_alignment": "center",
            "show_page_numbers": False,
        }, document_type="letter", logo=True)
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    sec = doc.sections[0]
    assert sec.page_width < sec.page_height
    assert sec.left_margin.inches > 1.1
    assert sec.top_margin.inches >= 1.12
    assert sec.footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert any("graphic" in p._p.xml for p in sec.header.paragraphs)
    assert "الموضوع:" in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize("bad_value,field", [
    ("invalid", "logo_alignment"),
    ("diagonal", "page_orientation"),
    ("giant", "page_margins"),
    (25, "body_font_size"),
])
def test_invalid_editor_choice_rejected(bad_value, field):
    with TestClient(app) as client:
        result = export(client, {field: bad_value})
    assert result.status_code == 422

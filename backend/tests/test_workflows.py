from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app

client = TestClient(app)
CONTENT = {
    'title': 'تحسين الخدمات', 'summary': 'خطة عمل قابلة للتنفيذ', 'language': 'ar',
    'sections': [
        {'title': 'المقدمة', 'body': ['تحديد احتياجات المستفيدين.'], 'notes': 'توضيح الهدف.'},
        {'title': 'التنفيذ', 'body': ['تبسيط الإجراءات.', 'قياس النتائج.'], 'notes': ''},
    ],
}


def test_docx_export_is_editable_and_arabic_rtl():
    response = client.post('/api/v1/exports/docx', json=CONTENT)
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    assert CONTENT['title'] in [p.text for p in doc.paragraphs]
    assert 'قياس النتائج.' in [p.text for p in doc.paragraphs]
    with ZipFile(BytesIO(response.content)) as archive:
        assert b'w:bidi' in archive.read('word/document.xml')


def test_pptx_export_keeps_text_and_speaker_notes():
    response = client.post('/api/v1/exports/pptx', json=CONTENT)
    assert response.status_code == 200
    deck = Presentation(BytesIO(response.content))
    assert len(deck.slides) == 3  # cover plus the two supplied sections
    assert any(CONTENT['title'] in s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert 'توضيح الهدف.' in deck.slides[1].notes_slide.notes_text_frame.text


def test_pptx_long_content_is_split_without_losing_text():
    content = {**CONTENT, 'sections': [{'title': 'المحور', 'body': ['تفاصيل ' * 70] * 4, 'notes': ''}]}
    response = client.post('/api/v1/exports/pptx', json=content)
    assert response.status_code == 200
    deck = Presentation(BytesIO(response.content))
    assert len(deck.slides) > 2
    actual = ' '.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert actual.count('تفاصيل') == 280


def test_analysis_does_not_invent_a_quality_score():
    doc = Document()
    doc.add_heading('عنوان', 1)
    doc.add_paragraph('نص تجريبي')
    output = BytesIO()
    doc.save(output)
    response = client.post('/api/v1/files/analyze', files={'file': ('sample.docx', output.getvalue())})
    assert response.status_code == 200
    assert response.json()['metrics']['headings'] == 1
    assert response.json()['health_score'] is None


def test_formatting_preserves_text_and_tables():
    doc = Document()
    doc.add_paragraph('المرجع 12345').add_run(' نص إضافي').bold = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'بيانات الجدول'
    output = BytesIO()
    doc.save(output)
    response = client.post('/api/v1/files/format-docx', files={'file': ('sample.docx', output.getvalue())})
    assert response.status_code == 200
    formatted = Document(BytesIO(response.content))
    assert formatted.paragraphs[0].text == 'المرجع 12345 نص إضافي'
    assert formatted.paragraphs[0].runs[1].bold is True
    assert formatted.tables[0].cell(0, 0).text == 'بيانات الجدول'


def test_unconfigured_ai_is_explicit(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    assert client.get('/api/v1/ai/status').json()['configured'] is False
    response = client.post('/api/v1/ai/generate', json={'topic': 'تحسين الخدمات الجامعية', 'section_count': 4})
    assert response.status_code == 503


def test_paid_generation_requires_access_token(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY', 'test-provider-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    monkeypatch.setenv('STUDIO_ACCESS_TOKEN', 'test-access-token')
    response = client.post('/api/v1/ai/generate', json={'topic': 'تحسين الخدمات الجامعية', 'section_count': 4})
    assert response.status_code == 401


def test_invalid_office_file_is_rejected_without_internal_details():
    response = client.post('/api/v1/files/analyze', files={'file': ('broken.docx', b'not a zip')})
    assert response.status_code == 422
    assert 'Traceback' not in response.text


def test_export_rejects_empty_content():
    response = client.post('/api/v1/exports/docx', json={**CONTENT, 'sections': []})
    assert response.status_code == 422


def test_provider_success_and_strict_schema(monkeypatch):
    import httpx
    import json
    from app import ai
    monkeypatch.setenv('OPENAI_API_KEY', 'test-provider-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    monkeypatch.setenv('STUDIO_ACCESS_TOKEN', 'test-access-token')
    original = httpx.Client
    def handle(request):
        payload = json.loads(request.content)
        assert payload['response_format']['json_schema']['strict'] is True
        assert request.headers['authorization'] == 'Bearer test-provider-key'
        return httpx.Response(200, json={'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(CONTENT)}}]})
    monkeypatch.setattr(ai.httpx, 'Client', lambda **kwargs: original(transport=httpx.MockTransport(handle), **kwargs))
    response = client.post('/api/v1/ai/generate', json={'topic': 'تحسين الخدمات', 'section_count': 2}, headers={'Authorization': 'Bearer test-access-token'})
    assert response.status_code == 200
    assert response.json() == CONTENT


def test_provider_error_is_not_exposed_and_no_fake_content(monkeypatch):
    import httpx
    from app import ai
    monkeypatch.setenv('OPENAI_API_KEY', 'test-provider-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    monkeypatch.setenv('STUDIO_ACCESS_TOKEN', 'test-access-token')
    original = httpx.Client
    monkeypatch.setattr(ai.httpx, 'Client', lambda **kwargs: original(transport=httpx.MockTransport(lambda r: httpx.Response(401, text='secret provider details')), **kwargs))
    response = client.post('/api/v1/ai/generate', json={'topic': 'تحسين الخدمات', 'section_count': 2}, headers={'Authorization': 'Bearer test-access-token'})
    assert response.status_code == 502
    assert 'secret' not in response.text
    assert 'sections' not in response.json()


def test_reference_mode_rejects_missing_source(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY', 'test-provider-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    monkeypatch.setenv('STUDIO_ACCESS_TOKEN', 'test-access-token')
    response = client.post('/api/v1/ai/generate', json={'topic': 'تحسين الخدمات', 'source_mode': 'reference'}, headers={'Authorization': 'Bearer test-access-token'})
    assert response.status_code == 422


def test_upload_limit(monkeypatch):
    from app import main
    monkeypatch.setattr(main, 'MAX_UPLOAD', 100)
    response = client.post('/api/v1/files/analyze', files={'file': ('large.docx', b'x' * 101)})
    assert response.status_code == 413


def test_ltr_export_does_not_force_rtl():
    content = {**CONTENT, 'language': 'en', 'title': 'Service improvement'}
    response = client.post('/api/v1/exports/docx', json=content)
    with ZipFile(BytesIO(response.content)) as archive:
        assert b'w:bidi w:val="0"' in archive.read('word/document.xml')


def test_rtl_alignment_uses_logical_start_for_word_and_libreoffice():
    response = client.post('/api/v1/exports/docx', json=CONTENT)
    with ZipFile(BytesIO(response.content)) as archive:
        xml = archive.read('word/document.xml')
        assert b'w:jc w:val="start"' in xml


def test_revision_reference_mode_rejects_missing_source(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY', 'test-provider-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    monkeypatch.setenv('STUDIO_ACCESS_TOKEN', 'test-access-token')
    response = client.post('/api/v1/ai/revise', json={'content': CONTENT, 'instruction': 'اختصر المحتوى', 'source_mode': 'reference'}, headers={'Authorization': 'Bearer test-access-token'})
    assert response.status_code == 422

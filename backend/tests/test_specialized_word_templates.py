"""Specialized report/study/strategy/letter DOCX templates preserve source text."""
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.ai import SYSTEM
from app.content import Content, GenerateRequest
from app.exports import export_docx
from app.main import app


def make_content():
    return Content(
        title="خطة تطوير إجراءات الخدمة",
        summary="تم إعداد عرض موجز للموضوع اعتمادًا على محتوى المستخدم.",
        language="ar",
        sections=[
            {
                "title": "الخلفية",
                "body": ["توضح الخلفية حدود العمل المطلوب دون إدخال معلومات مفترضة."],
                "notes": "",
            },
            {
                "title": "التوصيات",
                "body": ["تُدرس التوصيات بعد مراجعة المستندات والبيانات المعتمدة."],
                "notes": "",
            },
        ],
    )


def document_text(doc):
    values = [p.text for p in doc.paragraphs]
    values.extend(p.text for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs)
    return "\n".join(values)


@pytest.mark.parametrize(
    "kind,cover_label,special_heading",
    [
        ("general", "مستند احترافي", None),
        ("report", "تقرير رسمي", None),
        ("study", "دراسة تفصيلية", "خارطة محاور الدراسة"),
        ("strategy", "خطة استراتيجية", "محاور الخطة الاستراتيجية"),
    ],
)
def test_document_types_retain_user_text_and_native_bookmarked_toc(kind, cover_label, special_heading):
    content = make_content()
    data = export_docx(content, "luxury", kind)
    doc = Document(BytesIO(data))
    whole = document_text(doc)
    assert cover_label in whole
    assert "المحتويات" in whole
    for part in content.sections:
        assert part.title in whole
        for text in part.body:
            assert text in whole
    assert content.summary in whole
    if special_heading:
        assert special_heading in whole
        assert any(content.sections[0].title in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read("word/document.xml")
        assert b"w:bookmarkStart" in xml
        assert b"w:hyperlink" in xml
        assert b"w:bidi" in xml


@pytest.mark.parametrize("design", ["luxury", "corporate", "minimal"])
def test_formal_letter_is_compact_unbranded_and_does_not_invent_metadata(design):
    content = make_content()
    doc = Document(BytesIO(export_docx(content, design, "letter")))
    whole = document_text(doc)
    assert "الموضوع: " + content.title in whole
    assert content.summary in whole
    assert "المحتويات" not in whole
    assert "ملخص الدراسة" not in whole
    assert "MAGHRABI  /  OFFICE STUDIO" not in whole
    assert all(p.title in whole for p in content.sections)
    assert all(x in whole for section in content.sections for x in section.body)
    assert not doc.tables
    assert "MAGHRABI" not in doc.sections[0].header.paragraphs[0].text
    assert doc.sections[0].different_first_page_header_footer is False
    with ZipFile(BytesIO(export_docx(content, design, "letter"))) as archive:
        xml = archive.read("word/document.xml")
        assert b'w:bookmarkStart' not in xml
        assert b'w:hyperlink' not in xml


def test_specialized_word_export_endpoint_and_unknown_type_rejected():
    payload = make_content().model_dump()
    with TestClient(app) as client:
        response = client.post("/api/v1/exports/docx?design_style=corporate&document_type=letter", json=payload)
        assert response.status_code == 200
        doc = Document(BytesIO(response.content))
        assert "الموضوع:" in document_text(doc)
        invalid = client.post("/api/v1/exports/docx?document_type=unsupported", json=payload)
        assert invalid.status_code == 422
        legacy = client.post("/api/v1/exports/docx", json=payload)
        assert legacy.status_code == 200
        assert "المحتويات" in document_text(Document(BytesIO(legacy.content)))


def test_generate_request_supports_legacy_default_and_rejects_invalid_type():
    base = {"topic": "دراسة تنظيم الخدمات الإدارية"}
    assert GenerateRequest.model_validate(base).document_type == "general"
    assert GenerateRequest.model_validate({**base, "document_type": "letter"}).document_type == "letter"
    with pytest.raises(ValueError):
        GenerateRequest.model_validate({**base, "document_type": "fake"})
    assert "document_type" in SYSTEM

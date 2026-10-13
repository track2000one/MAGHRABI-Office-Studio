"""Word premium theme and native-editability contract checks."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from app.content import Content
from app.exports import export_docx
from app.main import app


def example(language="ar"):
    return Content(
        title="دراسة تطوير الخدمات الجامعية" if language == "ar" else "University services development",
        summary="تحسين تجربة المستفيدين وتطوير إجراءات العمل وقياس النتائج بصورة مستمرة.",
        language=language,
        sections=[
            {"title": "الأهداف", "body": ["تطوير جودة العمل.", "تحسين مستوى الخدمات.", "قياس نتائج الأداء."], "notes": ""},
            {"title": "الإجراءات التنفيذية", "body": ["حصر الوضع الراهن.", "إعداد خطة التنفيذ والمتابعة."], "notes": ""},
        ],
    )


def _all_text(document):
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        texts.extend(p.text for row in table.rows for cell in row.cells for p in cell.paragraphs)
    return texts


def test_premium_docx_has_cover_linked_contents_header_footer_and_editable_text():
    content = example()
    raw = export_docx(content)
    doc = Document(BytesIO(raw))
    contents = _all_text(doc)
    assert content.title in contents
    assert "المحتويات" in contents
    assert "الملخص التنفيذي" in contents
    for item in content.sections:
        assert item.title in contents
        for body in item.body:
            assert body in contents
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) >= 3
    assert all(item.title in [p.text for p in headings] for item in content.sections)
    assert doc.sections[0].different_first_page_header_footer is True
    assert len(doc.tables) >= 2
    assert doc.sections[0].page_width > doc.sections[0].page_height * 0.6
    with ZipFile(BytesIO(raw)) as archive:
        xml = archive.read("word/document.xml")
        assert b"w:bidi" in xml
        assert b"w:hyperlink" in xml
        assert b"w:bookmarkStart" in xml
        assert b"w:tbl" in xml
        footer = archive.read("word/footer1.xml")
        assert b"PAGE" in footer


def test_theme_choice_changes_accent_but_never_discards_text():
    content = example("en")
    colors = []
    for theme in ("luxury", "corporate", "minimal"):
        doc = Document(BytesIO(export_docx(content, theme)))
        assert content.title in _all_text(doc)
        assert all(x in _all_text(doc) for section in content.sections for x in section.body)
        colors.append(str(doc.styles["Heading 1"].font.color.rgb))
    assert len(set(colors)) == 3


def test_word_export_handles_long_arabic_body_without_truncation():
    body = " ".join(f"عنصر{i:03d}" for i in range(150))
    assert len(body) < 2000
    content = Content(title="الاستدامة", language="ar", sections=[
        {"title": "التحليل", "body": [body], "notes": ""},
    ])
    exported = Document(BytesIO(export_docx(content, "corporate")))
    assert body in [paragraph.text for paragraph in exported.paragraphs]


def test_word_endpoint_defaults_to_luxury_and_rejects_unknown_styles():
    payload = example().model_dump()
    with TestClient(app) as client:
        response = client.post("/api/v1/exports/docx", json=payload)
        assert response.status_code == 200
        assert "wordprocessingml.document" in response.headers["content-type"]
        assert Document(BytesIO(response.content)).paragraphs
        response = client.post("/api/v1/exports/docx?design_style=minimal", json=payload)
        assert response.status_code == 200
        invalid = client.post("/api/v1/exports/docx?design_style=unknown", json=payload)
        assert invalid.status_code == 422
